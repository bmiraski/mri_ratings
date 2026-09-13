"""Read Ben's historical MRI Excel workbooks (2003-2019) into normalized tables.

Every workbook in the Drive archive shares the same layout:

  Games      one row per game, team1/team2 with points, rushing, passing,
             turnovers, win flags, and an Excel date serial in column W
  Team Data  one row per FBS team (plus the pooled "Non D1A" row last),
             with the MRI in column V
  MRI        the published ranking table

Only the Games sheet is treated as source of truth. Everything on Team Data
and MRI is derived, and we recompute it in ratings.classic so the Python
implementation can be checked against the workbook's own numbers.
"""

from __future__ import annotations

import datetime as _dt
import re
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
import xlrd

POOLED_FCS = "Non D1A"

# Games sheet column indices (0-based).
_G = {
    "team1": 0,
    "team2": 1,
    "pts1": 2,
    "pts2": 3,
    "rush1": 4,
    "rush2": 5,
    "pass1": 6,
    "pass2": 7,
    "to1": 8,
    "to2": 9,
    "win1": 10,
    "win2": 11,
    "date": 22,
}

_NUMERIC = ["pts1", "pts2", "rush1", "rush2", "pass1", "pass2", "to1", "to2", "win1", "win2"]


@dataclass(frozen=True)
class ArchiveSeason:
    """One season parsed out of a workbook."""

    year: int
    path: Path
    games: pd.DataFrame
    teams: list[str]
    published: pd.DataFrame

    def __repr__(self) -> str:  # pragma: no cover - convenience only
        return (
            f"ArchiveSeason({self.year}, games={len(self.games)}, "
            f"teams={len(self.teams)}, published={len(self.published)})"
        )


def year_from_path(path: Path | str) -> int:
    """Pull the season year out of a filename.

    Football files name a single year (MRIFootball2010_final.xls). Basketball
    files name a span (MRIBasketball201920), and the season is labelled by the
    year it ends in - 2019-20 is season 2020 - which is also what the API uses.
    """
    name = Path(path).name
    # Basketball files named with a single year use the season's START year -
    # MRIBasketball2012 is the 2012-13 season, as its own title page says - so
    # it maps to season 2013 under the ends-in convention the API uses.
    single = re.fullmatch(r"MRIBasketball((?:19|20)\d{2})\.xlsx", name)
    if single:
        return int(single.group(1)) + 1
    span = re.search(r"((?:19|20)\d{2})(\d{2})(?!\d)", name)
    if span:
        start, end = span.group(1), span.group(2)
        return int(start[:2] + end) if end != start[2:] else int(start)
    match = re.search(r"(19|20)\d{2}", name)
    if not match:
        raise ValueError(f"no year in filename: {path}")
    return int(match.group(0))


def _num(cell, default=None):
    """Excel cells in these sheets are numbers, blanks, or stray strings."""
    value = cell.value
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        stripped = value.strip()
        if stripped:
            try:
                return float(stripped)
            except ValueError:
                return default
    return default


def _serial_to_date(serial: float | None, datemode: int) -> _dt.date | None:
    if serial is None or serial <= 0:
        return None
    try:
        return _dt.date(*xlrd.xldate_as_tuple(serial, datemode)[:3])
    except Exception:
        return None


def read_games(book: xlrd.Book) -> pd.DataFrame:
    """Normalize the Games sheet into one row per game."""
    sheet = book.sheet_by_name("Games")
    rows = []
    for r in range(1, sheet.nrows):
        row = sheet.row(r)
        team1 = str(row[_G["team1"]].value).strip()
        team2 = str(row[_G["team2"]].value).strip()
        if not team1 or not team2:
            continue

        record = {"team1": team1, "team2": team2}
        for key in _NUMERIC:
            record[key] = _num(row[_G[key]])

        # A row with no win flag on either side is a scheduled-but-unplayed
        # game; the workbook carries those as blanks.
        if record["win1"] is None and record["win2"] is None:
            continue
        record["win1"] = record["win1"] or 0.0
        record["win2"] = record["win2"] or 0.0

        serial = _num(row[_G["date"]]) if len(row) > _G["date"] else None
        record["date"] = _serial_to_date(serial, book.datemode)
        rows.append(record)

    frame = pd.DataFrame(rows)
    for key in _NUMERIC:
        frame[key] = frame[key].fillna(0.0).astype(float)
    return frame


def read_team_list(book: xlrd.Book) -> list[str]:
    """Team Data column A, in sheet order. The pooled FCS row sorts last."""
    sheet = book.sheet_by_name("Team Data")
    teams = []
    for r in range(1, sheet.nrows):
        name = str(sheet.cell(r, 0).value).strip()
        if not name:
            break
        teams.append(name)
    return teams


def read_published_ratings(book: xlrd.Book) -> pd.DataFrame:
    """The MRI sheet's published ranking, used as the validation target."""
    sheet = book.sheet_by_name("MRI")
    rows = []
    for r in range(2, sheet.nrows):
        team = str(sheet.cell(r, 1).value).strip()
        if not team:
            continue
        rating = _num(sheet.cell(r, 4))
        if rating is None:
            continue
        rows.append(
            {
                "rank": _num(sheet.cell(r, 0)),
                "team": team,
                "wins": _num(sheet.cell(r, 2), 0.0),
                "losses": _num(sheet.cell(r, 3), 0.0),
                "mri": rating,
                "mri_per_game": _num(sheet.cell(r, 5)),
            }
        )
    return pd.DataFrame(rows)


def read_workbook(path: Path | str) -> ArchiveSeason:
    path = Path(path)
    book = xlrd.open_workbook(path)
    return ArchiveSeason(
        year=year_from_path(path),
        path=path,
        games=read_games(book),
        teams=read_team_list(book),
        published=read_published_ratings(book),
    )


def read_archive(directory: Path | str) -> dict[int, ArchiveSeason]:
    """Read every workbook in a directory, keyed by season year.

    Where a year has several files (2010 ships a regular-season and a final
    workbook), the one with the most games wins.
    """
    seasons: dict[int, ArchiveSeason] = {}
    for path in sorted(Path(directory).glob("*.xls")):
        try:
            season = read_workbook(path)
        except Exception as exc:  # pragma: no cover - corrupt file guard
            print(f"  skipped {path.name}: {exc}")
            continue
        existing = seasons.get(season.year)
        if existing is None or len(season.games) > len(existing.games):
            seasons[season.year] = season
    return dict(sorted(seasons.items()))


# ---------------------------------------------------------------------------
# Basketball
# ---------------------------------------------------------------------------

BASKETBALL_GAMES = {
    "team1": 1, "team2": 2, "pts1": 3, "pts2": 4,
    "reb1": 5, "reb2": 6, "to1": 7, "to2": 8, "win1": 9, "win2": 10,
}
BASKETBALL_STATS = ["pts1", "pts2", "reb1", "reb2", "to1", "to2", "win1", "win2"]


def read_basketball_workbook(path: Path | str) -> ArchiveSeason:
    """Read one MRIBasketball workbook.

    Same idea as the football reader, three differences. The files are .xlsx so
    openpyxl reads them rather than xlrd; the statistics are rebounds and
    turnovers rather than rushing and passing; and there is no pooled opponent
    row, because the basketball workbooks drop non-D1 games instead of pooling
    them. The published table also starts on row 3, under a title row.
    """
    import openpyxl

    path = Path(path)
    book = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        games = _read_basketball_games(book["Games"])
        teams = _read_basketball_roster(book["Team Data"])
        published = _read_basketball_published(book["MRI"])
    finally:
        book.close()

    return ArchiveSeason(
        year=year_from_path(path),
        path=path,
        games=games,
        teams=teams,
        published=published,
    )


def _read_basketball_games(sheet) -> pd.DataFrame:
    rows = []
    for row in sheet.iter_rows(min_row=2, max_col=10, values_only=True):
        if not row or not row[0] or not row[1]:
            continue
        if not isinstance(row[2], (int, float)) or not isinstance(row[3], (int, float)):
            continue
        record = {"team1": str(row[0]).strip(), "team2": str(row[1]).strip()}
        for name, column in BASKETBALL_GAMES.items():
            if name in ("team1", "team2"):
                continue
            value = row[column - 1]
            record[name] = float(value) if isinstance(value, (int, float)) else 0.0
        rows.append(record)

    frame = pd.DataFrame(rows)
    if frame.empty:
        return frame
    for column in BASKETBALL_STATS:
        frame[column] = frame[column].astype(float)
    return frame


def _read_basketball_roster(sheet) -> list[str]:
    teams = []
    for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True):
        name = row[0]
        if not name:
            break
        teams.append(str(name).strip())
    return teams


def _read_basketball_published(sheet) -> pd.DataFrame:
    """The MRI sheet's ranking. Row 1 is a title, row 2 the header.

    Columns are located by header name rather than position, because the layout
    changed partway through: workbooks up to 2017-18 run Rank/Team/W/L/MRI,
    and from 2018-19 a Conference column is inserted third. Reading by position
    silently shifts W, L and MRI by one and yields ratings that look plausible
    and are wrong.
    """
    rows = list(sheet.iter_rows(min_row=1, max_row=2, max_col=12, values_only=True))
    header = rows[1] if len(rows) > 1 else ()
    index = {
        str(name).strip().casefold(): position
        for position, name in enumerate(header)
        if name
    }
    if "team" not in index or "mri" not in index:
        return pd.DataFrame()

    def value(row, key):
        position = index.get(key)
        return row[position] if position is not None and position < len(row) else None

    out = []
    for row in sheet.iter_rows(min_row=3, max_col=12, values_only=True):
        if not row or not value(row, "team"):
            continue
        rating = value(row, "mri")
        if not isinstance(rating, (int, float)):
            continue
        out.append(
            {
                "rank": value(row, "rank"),
                "team": str(value(row, "team")).strip(),
                "conference": value(row, "conference"),
                "wins": value(row, "w"),
                "losses": value(row, "l"),
                "mri": float(rating),
            }
        )
    return pd.DataFrame(out)


MIN_USABLE_GAMES = 500


def read_basketball_archive(directory: Path | str) -> dict[int, ArchiveSeason]:
    """Read every usable basketball workbook in a directory.

    Some files in the archive are stubs rather than seasons - MRIBasketball2013
    carries 163 games and a published table that is entirely #DIV/0!, because it
    was saved before the season had begun. Validating against one proves
    nothing, so anything without a real game log and a real published table is
    skipped rather than silently failed.
    """
    seasons: dict[int, ArchiveSeason] = {}
    for path in sorted(Path(directory).glob("MRIBasketball*.xlsx")):
        try:
            season = read_basketball_workbook(path)
        except Exception as exc:  # pragma: no cover - corrupt file guard
            print(f"  skipped {path.name}: {exc}")
            continue
        if len(season.games) < MIN_USABLE_GAMES or season.published.empty:
            print(f"  skipped {path.name}: {len(season.games)} games, "
                  f"{len(season.published)} published ratings - not a usable season")
            continue
        seasons[season.year] = season
    return dict(sorted(seasons.items()))
