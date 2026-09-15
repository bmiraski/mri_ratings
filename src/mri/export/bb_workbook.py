"""Write a basketball season out as an Excel workbook in the shape of the
original 2003-2019 workbooks - see ``export/workbook.py`` for the football
version this mirrors and for the reasoning behind writing formulas instead of
values.

Basketball differs from football in four places that matter (see
``ratings/classic.py`` for the authoritative statement): the margin cap is 30
not 35, there is no undefeated-opponent fallback, the statistical components
are rebound and turnover differential per game rather than rush/pass/defense,
and strength of schedule is RPI-style - it strips a team's own games out of
its opponents' records before computing OppWinPct and OppOppWinPct for that
purpose. Every one of those differences has to survive the trip into Excel
formulas, or the sheet quietly rates a different game than the Python does.

The other structural difference is bigger than any formula: football pools
every non-FBS opponent into one "Non D1A" row, so its Team Data sheet only
ever excludes a single team from the rated field. Basketball has no pooled
team at all - a non-D1 opponent (a D2 or NAIA program that a D1 team
scheduled out of conference) keeps its own name and its own row, so that
Games-sheet lookups resolve, but it is left out of the MRI ranking and out of
the z-score mean/stdev individually, the same way ``classic.compute`` leaves
it out of the rated field it returns. There can be dozens of these a season,
not one.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from ..ingest import bb_registry as registry
from .bb_sitedata import season_label

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")

GAMES_HEADERS = [
    "Team1", "Team2", "Pnt1", "Pnt2", "Reb1", "Reb2", "TO1", "TO2", "WT1", "WT2",
    "T1W", "T1L", "T2W", "T2L", "T1OW", "T1OL", "T2OW", "T2OL", "T1Point", "T2Point",
    "Date",
]

# The 20 labelled columns match the 2019-20 workbook's Team Data sheet exactly.
# SOS and SOS Rank are new: Ben's original computed the same RPI arithmetic but
# left it in unlabelled columns (his Y and Z) past the labelled set, so it was
# there without being named. We label it, and - following his own layout -
# still keep the two raw pieces (opponents' win pct and opponents' opponents'
# win pct, each with the team's own games backed out) as unlabelled working
# columns past SOS Rank, because the SOS formula reads more clearly as a
# product of two named cells than as one formula with the RPI arithmetic
# inlined twice.
TEAM_HEADERS = [
    "Team", "W", "L", "Opp W", "Opp L", "O2 W", "O2 L", "Pnts", "Rebs", "O Rebs",
    "TO's", "O TO's", "Away Win", "Away Loss", "Reb Diff", "RDPG", "TO Diff", "TODPG",
    "Conference", "MRI Score", "SOS", "SOS Rank",
]


def _style_header(sheet, row: int, count: int) -> None:
    for column in range(1, count + 1):
        cell = sheet.cell(row, column)
        cell.font = Font(name=FONT, bold=True, size=10)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def write_season(
    season: int,
    games: pd.DataFrame,
    modern: pd.DataFrame | None,
    path: Path,
) -> Path:
    """Build one season's workbook.

    ``games`` carries ``cbbd.classic_table``'s raw column names and raw team
    spellings - it is canonicalized here, the same way ``bb_sitedata._canonical``
    does it, rather than upstream, because the rated field (which teams are D1
    for *this* season) has to be decided against the same season the canonical
    spelling is resolved against.
    """
    games = games.copy()
    for column in ("team1", "team2"):
        games[column] = [registry.resolve(n, n, season=season) for n in games[column]]

    all_teams = sorted(set(games["team1"]) | set(games["team2"]))
    d1 = [t for t in all_teams if registry.is_d1(t, season=season)]
    non_d1 = [t for t in all_teams if t not in set(d1)]
    ordered = d1 + non_d1

    book = Workbook()
    book.remove(book.active)

    _about_sheet(book, season, len(games), len(d1),
                 int(games["reb1"].isna().sum()))
    _games_sheet(book, games, len(ordered))
    stats_row = _team_data_sheet(book, games, ordered, d1, season)
    _mri_sheet(book, d1, len(ordered))
    if modern is not None and not modern.empty:
        _modern_sheet(book, modern, season)

    _define_names(book, stats_row)

    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path


def _about_sheet(book: Workbook, season: int, game_count: int, d1_count: int,
                 no_box: int = 0) -> None:
    sheet = book.create_sheet("About")
    gap = []
    if no_box:
        gap = [
            ("MISSING BOX SCORES", True),
            (f"{no_box} of these {game_count} games have a score but no rebound or", False),
            ("turnover counts in the source data. Their Reb and TO cells are blank", False),
            ("rather than zero, and column Y of Team Data counts the games that do", False),
            ("have counts - that is what RDPG and TODPG divide by, so a missing box", False),
            ("score costs a team nothing but the statistics it never reported.", False),
            ("", False),
        ]
    lines = [
        (f"MRI {season_label(season)}", True),
        ("", False),
        (f"{game_count} games, {d1_count} D1 teams.", False),
        ("", False),
        *gap,
        ("WHAT IS LIVE", True),
        ("Games, Team Data and MRI carry the original formulas. Correct a score on", False),
        ("the Games sheet and the whole season re-rates, exactly as the 2003-2019", False),
        ("workbooks did. Only columns A-J and Date are data; K through T are formulas.", False),
        ("", False),
        ("WHAT IS A SNAPSHOT", True),
        ("The MRI 2.0 sheet is values, not formulas. That rating solves every game", False),
        ("in the season at once as a regularized least-squares system, which a", False),
        ("spreadsheet cannot express. Editing a score will not update it; rerun", False),
        ("scripts/export_bb_workbooks.py instead.", False),
        ("", False),
        ("NON-D1 OPPONENTS", True),
        ("Basketball keeps no pooled team, unlike football's \"Non D1A\". A non-D1", False),
        ("opponent appears under its own name, so Games-sheet lookups still resolve,", False),
        ("but it is simply left out of the MRI ranking and out of the z-score mean", False),
        ("and standard deviation - excluded individually, not merged into anything.", False),
        ("", False),
        ("Generated by github.com/bmiraski/mri_ratings", False),
    ]
    for index, (text, bold) in enumerate(lines, start=1):
        cell = sheet.cell(index, 1, text)
        cell.font = Font(name=FONT, bold=bold, size=13 if index == 1 else 10)
    sheet.column_dimensions["A"].width = 78
    # By heading rather than by row number: the missing-box-score section is
    # only present in some seasons, so fixed row indices highlighted the wrong
    # lines in the seasons that have it.
    highlight = {"WHAT IS LIVE", "MISSING BOX SCORES", "NON-D1 OPPONENTS"}
    for index, (text, _) in enumerate(lines, start=1):
        if text in highlight:
            sheet.cell(index, 1).fill = NOTE_FILL


def _games_sheet(book: Workbook, games: pd.DataFrame, team_rows: int) -> None:
    sheet = book.create_sheet("Games")
    sheet.append(GAMES_HEADERS)
    _style_header(sheet, 1, len(GAMES_HEADERS))

    last = team_rows + 1  # Team Data lookup range ends here
    # Only columns 2-5 (W, L, Opp W, Opp L) are ever looked up, but the range is
    # given generously the way the football sheet's is.
    table = f"'Team Data'!$A$1:$E${last}"

    for offset, row in enumerate(games.itertuples(), start=2):
        sheet.cell(offset, 1, row.team1)
        sheet.cell(offset, 2, row.team2)
        for column, value in enumerate(
            [row.pts1, row.pts2, row.reb1, row.reb2, row.to1, row.to2, row.win1, row.win2],
            start=3,
        ):
            # A game whose box score never arrived leaves Reb and TO blank
            # rather than zero, which is what lets Team Data count the games
            # that actually have statistics. Scores and win flags are never
            # missing - the games feed carries those for every final.
            if pd.notna(value):
                sheet.cell(offset, column, float(value))
        if pd.notna(row.start_date):
            # openpyxl rejects tz-aware datetimes outright, and the time of
            # day carries no information the ratings use, so only the
            # calendar date survives into the cell.
            sheet.cell(offset, 21, pd.to_datetime(row.start_date, utc=True).date())

        r = offset
        # Opponent records, looked up the way the originals did.
        sheet.cell(r, 11, f"=VLOOKUP(A{r},{table},2,FALSE)")
        sheet.cell(r, 12, f"=VLOOKUP(A{r},{table},3,FALSE)")
        sheet.cell(r, 13, f"=VLOOKUP(B{r},{table},2,FALSE)")
        sheet.cell(r, 14, f"=VLOOKUP(B{r},{table},3,FALSE)")
        sheet.cell(r, 15, f"=VLOOKUP(A{r},{table},4,FALSE)")
        sheet.cell(r, 16, f"=VLOOKUP(A{r},{table},5,FALSE)")
        sheet.cell(r, 17, f"=VLOOKUP(B{r},{table},4,FALSE)")
        sheet.cell(r, 18, f"=VLOOKUP(B{r},{table},5,FALSE)")
        # Game credit: capped margin (30, not football's 35) weighted by how
        # good the opponent was. Basketball has no undefeated-opponent
        # fallback, so the loss branch's opponent-loss-pct term is a plain
        # ratio with no IF(...,0.1,...) substitution - copying football's
        # fallback in would change every rating in a season where anyone runs
        # the table.
        sheet.cell(r, 19, (
            f"=IF(I{r}=1,(MIN(30,(C{r}-D{r})))*(M{r}/(M{r}+N{r}))*(Q{r}/(Q{r}+R{r})),"
            f"(MAX(-30,(C{r}-D{r})))*(N{r}/(M{r}+N{r}))*(R{r}/(R{r}+Q{r})))"
        ))
        sheet.cell(r, 20, (
            f"=IF(J{r}=1,(MIN(30,(D{r}-C{r})))*(K{r}/(K{r}+L{r}))*(O{r}/(O{r}+P{r})),"
            f"(MAX(-30,(D{r}-C{r})))*(L{r}/(K{r}+L{r}))*(P{r}/(P{r}+O{r})))"
        ))

    for column in range(1, len(GAMES_HEADERS) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 18 if column < 3 else 9
    sheet.column_dimensions["U"].width = 12
    sheet.freeze_panes = "C2"


def _team_data_sheet(
    book: Workbook,
    games: pd.DataFrame,
    ordered: list[str],
    d1: list[str],
    season: int,
) -> int:
    sheet = book.create_sheet("Team Data")
    sheet.append(TEAM_HEADERS)
    _style_header(sheet, 1, len(TEAM_HEADERS))

    n = len(games)
    g = f"Games!$A$2:$A${n + 1}"
    h = f"Games!$B$2:$B${n + 1}"

    def both(column: str, away_column: str) -> str:
        """SUMIF over both sides of the game log, as the originals did."""
        return (f"SUMIF({g},$A{{r}},Games!{column}$2:{column}${n + 1})"
                f"+SUMIF({h},$A{{r}},Games!{away_column}$2:{away_column}${n + 1})")

    d1_last = len(d1) + 1  # last D1 row: D1 teams sort first, non-D1 after
    last_team_row = len(ordered) + 1

    for offset, team in enumerate(ordered, start=2):
        r = offset
        rated = team in set(d1)
        sheet.cell(r, 1, team)
        sheet.cell(r, 2, "=" + both("$I", "$J").format(r=r))    # W
        # L is played-minus-won, not "sum of the opponent's win flag" - the
        # two agree for every ordinary game (win1+win2 == 1) but not for a
        # cancelled game the feed still records as a 0-0 final, where both
        # win flags are 0. classic.py's own losses are built the same way -
        # long["lost"] = 1 - long["won"], summed - so a 0-0 row counts as a
        # loss for BOTH sides there, not as a no-decision for either; summing
        # the opponent's flag instead undercounted exactly that case and was
        # the one place a real season (2021-22) failed the recalculation
        # check.
        sheet.cell(r, 3, (
            f"=COUNTIF(Games!$A:$A,$A{r})+COUNTIF(Games!$B:$B,$A{r})-B{r}"
        ))
        sheet.cell(r, 4, "=" + both("$M", "$K").format(r=r))    # Opp W
        sheet.cell(r, 5, "=" + both("$N", "$L").format(r=r))    # Opp L
        sheet.cell(r, 6, "=" + both("$Q", "$O").format(r=r))    # O2 W
        sheet.cell(r, 7, "=" + both("$R", "$P").format(r=r))    # O2 L
        sheet.cell(r, 8, "=" + both("$S", "$T").format(r=r))    # Pnts (game credit)
        sheet.cell(r, 9, "=" + both("$E", "$F").format(r=r))    # Rebs
        sheet.cell(r, 10, "=" + both("$F", "$E").format(r=r))   # O Rebs
        sheet.cell(r, 11, "=" + both("$G", "$H").format(r=r))   # TO's
        sheet.cell(r, 12, "=" + both("$H", "$G").format(r=r))   # O TO's
        # Away Win/Loss: team1 is the visitor, so only the team1 side counts.
        # Loss is away-appearances-minus-away-wins for the same reason the
        # overall L column is: it must not undercount a 0-0 cancelled game.
        sheet.cell(r, 13, f"=SUMIF(Games!$A:$A,$A{r},Games!I:I)")
        sheet.cell(r, 14, f"=COUNTIF(Games!$A:$A,$A{r})-M{r}")
        # A non-D1 opponent needs the raw SUMIF totals above so that a D1
        # team's own lookups resolve, but nothing from here on is meaningful
        # for a team that is not part of the rated field: it gets no
        # per-game rate, no conference, no rating and no schedule strength -
        # left out individually, the same way classic.compute's z() only
        # pools the rated teams for its mean and standard deviation.
        if not rated:
            continue
        sheet.cell(r, 15, f"=I{r}-J{r}")                                    # Reb Diff
        sheet.cell(r, 16, f"=IFERROR(O{r}/Y{r},0)")                          # RDPG
        sheet.cell(r, 17, f"=L{r}-K{r}")                                    # TO Diff
        sheet.cell(r, 18, f"=IFERROR(Q{r}/Y{r},0)")                          # TODPG
        sheet.cell(r, 19, registry.conference_of(team, season=season) or "")
        sheet.cell(r, 20, (
            f"=IFERROR((B{r}/(B{r}+C{r}))*25,0)+IFERROR((D{r}/(D{r}+E{r}))*25,0)"
            f"+IFERROR((F{r}/(F{r}+G{r}))*10,0)+H{r}"
            f"+((P{r}-RDAVG)/RebSD)*10+((R{r}-TODAVG)/TOSD)*6"
        ))
        # SOS: RPI-style, backing the team's own games out of its opponents'
        # combined record (see classic._sos_basketball). W and X are the two
        # working pieces - opponents' win pct and opponents' opponents' win
        # pct, each with the team's own games removed - kept as separate
        # unlabelled columns past SOS Rank so the SOS formula itself is just
        # their product, matching how the 2019-20 workbook laid it out.
        sheet.cell(r, 21, f"=W{r}*X{r}")
        sheet.cell(r, 22, f"=RANK(U{r},U$2:U${d1_last})")
        sheet.cell(r, 23, f"=IFERROR((D{r}-C{r})/(D{r}+E{r}-(B{r}+C{r})),0)")
        sheet.cell(r, 24, f"=IFERROR((F{r}-((B{r}+C{r})*B{r}))/((F{r}+G{r})-(B{r}+C{r})^2),0)")
        # Y: games with a box score, which is the denominator RDPG and TODPG
        # divide by. It equals W+L wherever the feed is complete, and in the
        # two seasons where it is not - 2004-05 and 2011-12, missing about one
        # box score in eight - dividing by W+L would shrink a team's rebound
        # margin in proportion to how many of its box scores are missing. Kept
        # past the labelled set, beside the other working columns.
        sheet.cell(r, 25, (
            f"=COUNTIFS({g},$A{r},Games!$E$2:$E${n + 1},\"<>\")"
            f"+COUNTIFS({h},$A{r},Games!$F$2:$F${n + 1},\"<>\")"
        ))

    stats_row = last_team_row + 1
    sheet.cell(stats_row, 1, "Mean").font = Font(name=FONT, bold=True, size=10)
    sheet.cell(stats_row + 1, 1, "StDev").font = Font(name=FONT, bold=True, size=10)
    for column in (16, 18):  # RDPG, TODPG
        letter = get_column_letter(column)
        sheet.cell(stats_row, column, f"=AVERAGE({letter}2:{letter}{d1_last})")
        sheet.cell(stats_row + 1, column, f"=STDEV({letter}2:{letter}{d1_last})")

    note = sheet.cell(stats_row + 3, 1,
                      "Mean and StDev cover D1 teams only (rows 2 through "
                      f"{d1_last}); non-D1 opponents sort after and are excluded, "
                      "as in classic.compute's z-score pool.")
    note.font = Font(name=FONT, italic=True, size=9)

    sheet.column_dimensions["A"].width = 24
    for column in range(2, len(TEAM_HEADERS) + 3):
        sheet.column_dimensions[get_column_letter(column)].width = 10
    sheet.freeze_panes = "B2"
    return stats_row


def _mri_sheet(book: Workbook, d1: list[str], team_rows: int) -> None:
    sheet = book.create_sheet("MRI", 0)
    sheet.append(["Rank", "Team", "Conference", "W", "L", "MRI"])
    _style_header(sheet, 1, 6)

    last = team_rows + 1
    table = f"'Team Data'!$A$2:$T${last}"

    for offset, team in enumerate(d1, start=2):
        r = offset
        sheet.cell(r, 1, f"=RANK(F{r},F$2:F${len(d1) + 1})")
        sheet.cell(r, 2, team)
        sheet.cell(r, 3, f"=VLOOKUP(B{r},{table},19,FALSE)")
        sheet.cell(r, 4, f"=VLOOKUP(B{r},{table},2,FALSE)")
        sheet.cell(r, 5, f"=VLOOKUP(B{r},{table},3,FALSE)")
        sheet.cell(r, 6, f"=VLOOKUP(B{r},{table},20,FALSE)")
        sheet.cell(r, 6).number_format = "0.000"

    sheet.column_dimensions["B"].width = 24
    sheet.column_dimensions["C"].width = 18
    for column in ("A", "D", "E", "F"):
        sheet.column_dimensions[column].width = 10
    sheet.freeze_panes = "A2"


def _modern_sheet(book: Workbook, modern: pd.DataFrame, season: int) -> None:
    sheet = book.create_sheet("MRI 2.0", 1)
    sheet.append(["Rank", "Team", "Conference", "Power", "Resume", "Resume Rank", "Games"])
    _style_header(sheet, 1, 7)

    for offset, row in enumerate(modern.itertuples(), start=2):
        sheet.cell(offset, 1, int(row.rank))
        sheet.cell(offset, 2, row.team)
        sheet.cell(offset, 3, registry.conference_of(row.team, season=season) or "")
        sheet.cell(offset, 4, round(float(row.power), 2)).number_format = "+0.00;-0.00"
        if pd.notna(row.resume):
            sheet.cell(offset, 5, round(float(row.resume), 2)).number_format = "+0.00;-0.00"
            sheet.cell(offset, 6, int(row.resume_rank))
        sheet.cell(offset, 7, int(row.games))

    note = sheet.cell(len(modern) + 3, 1,
                      "Snapshot, not formulas: MRI 2.0 solves the whole season at once as a "
                      "regularized least-squares system, which a spreadsheet cannot express. "
                      "Power is points against an average D1 team; Resume is wins above what "
                      "an average team would manage against the same schedule.")
    note.font = Font(name=FONT, italic=True, size=9)
    note.fill = NOTE_FILL

    sheet.column_dimensions["B"].width = 24
    sheet.column_dimensions["C"].width = 18
    for column in ("A", "D", "E", "F", "G"):
        sheet.column_dimensions[column].width = 12
    sheet.freeze_panes = "A2"


def _define_names(book: Workbook, stats_row: int) -> None:
    """The named ranges the MRI formula reads, as in the originals."""
    names = {
        "RDAVG": f"P${stats_row}", "RebSD": f"P${stats_row + 1}",
        "TODAVG": f"R${stats_row}", "TOSD": f"R${stats_row + 1}",
    }
    for name, reference in names.items():
        book.defined_names[name] = DefinedName(name, attr_text=f"'Team Data'!${reference}")
